from __future__ import annotations

import gzip
import json
from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session, sessionmaker
from starlette.middleware.sessions import SessionMiddleware

import jobradar.main as main_module
from jobradar.config import Settings
from jobradar.db import Base, build_engine, get_db
from jobradar.integrations.codex.provider import CodexJobSummary, GeneratedRoleTermGroup
from jobradar.models import (
    ApplicationDraft,
    ApplicationRecord,
    AppSetting,
    CandidateAddress,
    CandidateProfile,
    CrawlRun,
    JobPosting,
    JobScore,
    JobSnapshot,
    PreferenceProfile,
    Source,
)
from jobradar.services import settings as setting_service
from jobradar.services.application_drafts import (
    ApplicationDraftReview,
    DraftClaim,
    DraftRequest,
    GeneratedApplicationDraft,
)
from jobradar.web import router


class WebDraftProvider:
    async def generate_application_draft(
        self, request: DraftRequest
    ) -> GeneratedApplicationDraft:
        return GeneratedApplicationDraft(
            cv_draft="Lebenslauf mit Kalendersteuerung",
            cover_letter="Motivationsschreiben für die Geschäftsführung",
            evidence_map=(
                DraftClaim(
                    claim="Kalendersteuerung",
                    evidence_ids=(request.facts[0].evidence_id,),
                ),
            ),
        )

    async def review_application_draft(
        self,
        request: DraftRequest,
        draft: GeneratedApplicationDraft,
    ) -> ApplicationDraftReview:
        del request, draft
        return ApplicationDraftReview(
            decision="approve",
            scores={"relevance": 5, "evidence": 5, "clarity": 4, "motivation": 4},
            must_fix=(),
            optional_improvements=(),
            unsupported_claims=(),
            summary="Freigegeben.",
        )


@pytest.fixture
def session_factory() -> sessionmaker[Session]:
    engine = build_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False, class_=Session)
    yield factory
    engine.dispose()


@pytest.fixture
def client(session_factory: sessionmaker[Session]) -> TestClient:
    app = FastAPI()
    app.state.settings = Settings(default_locale="de")
    app.add_middleware(SessionMiddleware, secret_key="test-session-secret")
    app.include_router(router)

    def override_db():
        session = session_factory()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_db
    with TestClient(app) as test_client:
        yield test_client


def _seed_jobs(session_factory: sessionmaker[Session]) -> list[int]:
    now = datetime.now(UTC)
    with session_factory() as session:
        source = Source(
            name="Bundesagentur für Arbeit",
            slug="arbeitsagentur",
            kind="api",
            enabled=True,
            is_official=False,
            free_api=True,
            status="ready",
        )
        session.add(source)
        jobs = []
        values = (
            ("Assistenz der Geschäftsführung", "Muster GmbH", "Stuttgart", "onsite", 91, 84),
            ("Executive Assistant", "Beispiel AG", "Karlsruhe", "hybrid", 82, 76),
            ("CEO Office Coordinator", "Remote SE", "Deutschland", "remote", 74, 69),
        )
        for index, (title, employer, location, remote_type, relevance, fit) in enumerate(values):
            job = JobPosting(
                source=source,
                external_id=f"job-{index}",
                canonical_url=f"https://example.invalid/jobs/{index}",
                title=title,
                employer=employer,
                description_text=(
                    "Du koordinierst Termine und bereitest Entscheidungen vor.\n\n"
                    "Deine Aufgaben:\n"
                    "- Termine und Reisen koordinieren\n"
                    "- Entscheidungen vorbereiten\n\n"
                    "<script>alert('nicht ausführen')</script>"
                ),
                location_text=location,
                state="Baden-Württemberg" if remote_type != "remote" else None,
                remote_type=remote_type,
                employment_type="Vollzeit",
                contract_type="Unbefristet",
                published_at=now,
                first_seen_at=now,
                last_seen_at=now,
                is_active=True,
            )
            job.scores.append(
                JobScore(
                    relevance_score=relevance,
                    candidate_fit_score=fit,
                    relevance_reasons=[{"reason": "Direkte Unterstützung der Geschäftsführung"}],
                    fit_reasons=[{"reason": "Passt zu deiner Organisationserfahrung"}],
                )
            )
            jobs.append(job)
        jobs[0].application = ApplicationRecord(status="Merkliste")
        session.add_all(jobs)
        session.add(
            CandidateProfile(
                full_name="Anna Beispiel",
                structured_data={"current_title": "Teamassistenz", "experience_years": 5},
            )
        )
        session.add(
            PreferenceProfile(
                employment_types=["Vollzeit"],
                contract_types=["Unbefristet"],
                preferred_states=["Baden-Württemberg"],
                include_germany_remote=True,
                languages=["Deutsch", "Englisch"],
                primary_role_terms=["Assistenz der Geschäftsführung"],
                additional_role_terms=[],
            )
        )
        session.commit()
        return [job.id for job in jobs]


def _cv_docx() -> bytes:
    document = Document()
    document.add_paragraph("Anna Beispiel")
    document.add_paragraph("Work Experience")
    document.add_paragraph("2020 - 2025 Office Manager at Beispiel AG")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_dashboard_is_directly_accessible(client: TestClient) -> None:
    dashboard = client.get("/")
    assert dashboard.status_code == 200
    assert "Jobportale: nicht geprüft" in dashboard.text
    assert "Firmenportale: nicht geprüft" in dashboard.text
    assert "APIs: nicht geprüft" in dashboard.text
    assert "ChatGPT: nicht geprüft" in dashboard.text
    assert "Betriebsstatus" not in dashboard.text
    assert 'data-system-chatgpt' in dashboard.text
    assert 'data-active-label="ChatGPT verbunden"' in dashboard.text
    assert "data-system-job-portals" in dashboard.text
    assert "data-system-company-sites" in dashboard.text
    assert 'data-running-label="Jobportale: Crawl läuft"' in dashboard.text
    assert 'data-found-label="gefunden"' in dashboard.text
    assert 'data-codex-summary' not in dashboard.text
    assert "OpenAI wird geprüft; Crawling bleibt deaktiviert." not in dashboard.text
    assert "Sicherer Grundbetrieb aktiv" not in dashboard.text
    assert "Anwendung bereit" not in dashboard.text
    assert "Privater Zugang eingerichtet" not in dashboard.text
    assert "Stellensuche starten" not in dashboard.text

    status_script = client.get("/static/app.js")
    assert status_script.status_code == 200
    assert "getJson('/api/codex/status')" in status_script.text
    assert "data-codex-step" not in status_script.text
    assert "data-codex-summary" not in status_script.text


def test_application_factory_exposes_dashboard_directly(
    session_factory: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(main_module, "upgrade_database", lambda **_kwargs: None)
    monkeypatch.setattr(main_module, "get_session_factory", lambda: session_factory)
    app = main_module.create_app(
        Settings(database_url="sqlite://", secret_key="test-session-secret-at-least-32-bytes")
    )

    def override_db():
        session = session_factory()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_db
    with TestClient(app) as application_client:
        dashboard = application_client.get("/", follow_redirects=False)
        assert dashboard.status_code == 200


def test_jobs_are_split_by_workplace_and_filterable(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _seed_jobs(session_factory)

    response = client.get("/jobs")
    assert response.status_code == 200
    assert "Präsenz in Baden-Württemberg" in response.text
    assert "Hybrid in Baden-Württemberg" in response.text
    assert "Deutschlandweit Remote" in response.text
    assert "Assistenz der Geschäftsführung" in response.text
    assert "91 %" in response.text
    assert "Arbeitgeber" in response.text
    assert "Muster GmbH" in response.text
    assert 'class="job-excerpt"' in response.text
    assert 'name="workplace"' not in response.text
    assert 'name="employment"' not in response.text
    assert 'name="contract"' not in response.text

    filtered = client.get("/jobs", params={"min_score": 80})
    assert filtered.status_code == 200
    assert "Assistenz der Geschäftsführung" in filtered.text
    assert "Executive Assistant" in filtered.text
    assert "CEO Office Coordinator" not in filtered.text


def test_jobs_default_does_not_reapply_removed_settings_filters(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    job_ids = _seed_jobs(session_factory)
    with session_factory() as session:
        job = session.get(JobPosting, job_ids[1])
        assert job is not None
        job.employment_type = "Teilzeit"
        job.contract_type = "Befristet"
        session.add_all(
            [
                AppSetting(key="ui.fulltime_default", value_json=True),
                AppSetting(key="ui.permanent_default", value_json=True),
            ]
        )
        session.commit()

    response = client.get("/jobs")

    assert response.status_code == 200
    assert "Executive Assistant" in response.text
    assert "Teilzeit" in response.text
    assert "Befristet" in response.text


def test_title_matches_include_low_relevance_jobs_and_exclude_description_only_matches(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    job_ids = _seed_jobs(session_factory)
    with session_factory() as session:
        title_match = session.get(JobPosting, job_ids[1])
        assert title_match is not None
        title_match.scores[0].relevance_score = 15
        title_match.scores[0].created_at = datetime.now(UTC)
        title_match.remote_type = "onsite"
        description_only = session.get(JobPosting, job_ids[2])
        assert description_only is not None
        description_only.title = "Operations Coordinator"
        description_only.description_text += " Executive Assistant"
        preference = session.query(PreferenceProfile).one()
        preference.additional_role_terms = [
            *preference.additional_role_terms,
            "Executive Assistant",
        ]
        session.commit()

    response = client.get("/jobs", params={"view": "title_matches"})

    assert response.status_code == 200
    assert "Passende Stellentitel" in response.text
    assert "Executive Assistant" in response.text
    assert "15 %" in response.text
    assert "Operations Coordinator" not in response.text


def test_jobs_render_at_most_100_per_section_and_load_the_next_batch(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    now = datetime.now(UTC)
    with session_factory() as session:
        source = Source(
            name="Testquelle",
            slug="batch-source",
            kind="api",
            enabled=True,
            is_official=False,
            free_api=True,
            status="ready",
        )
        session.add(source)
        for index in range(105):
            job = JobPosting(
                source=source,
                external_id=f"batch-{index}",
                canonical_url=f"https://example.invalid/batch/{index}",
                title=f"Executive Assistant Batch {index:03}",
                employer="Batch GmbH",
                description_text="Unterstützung der Geschäftsführung",
                location_text="Stuttgart",
                state="Baden-Württemberg",
                remote_type="onsite",
                first_seen_at=now,
                last_seen_at=now,
                is_active=True,
            )
            job.scores.append(JobScore(relevance_score=15, candidate_fit_score=20))
            session.add(job)
        session.commit()

    first_page = client.get("/jobs")
    assert first_page.status_code == 200
    assert first_page.text.count('class="job-card ') == 100
    assert ">105<" in first_page.text
    assert "Executive Assistant Batch 004" not in first_page.text
    assert "data-job-list-sentinel" in first_page.text

    next_page = client.get(
        "/jobs/more",
        params={"section": "onsite", "offset": 100, "view": "all", "min_score": 0},
    )
    assert next_page.status_code == 200
    assert next_page.text.count('class="job-card ') == 5
    assert "Executive Assistant Batch 004" in next_page.text
    assert next_page.headers["X-Next-Offset"] == "105"
    assert next_page.headers["X-Has-More"] == "false"


def test_job_detail_shows_both_scores_reasons_and_source(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    job_id = _seed_jobs(session_factory)[0]
    response = client.get(f"/jobs/{job_id}")

    assert response.status_code == 200
    assert "Aufgabennähe" in response.text
    assert "Kandidatenpassung" in response.text
    assert "Direkte Unterstützung der Geschäftsführung" in response.text
    assert "Bundesagentur für Arbeit" in response.text
    assert "Muster GmbH" in response.text
    assert "Stellendaten" in response.text
    assert 'class="prose-body"' in response.text
    assert "<h3>Deine Aufgaben</h3>" in response.text
    assert "<li>Termine und Reisen koordinieren</li>" in response.text
    assert "<script>alert" not in response.text
    assert "&lt;script&gt;" in response.text


def test_job_detail_generates_and_caches_luna_summary(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    job_id = _seed_jobs(session_factory)[0]
    raw_description = (
        "<h2>Deine Aufgaben</h2><p>Du unterstützt die Geschäftsführung.</p>"
        "<ul><li>Termine koordinieren</li><li>Reisen planen</li></ul>"
        "<script>alert('nicht ausführen')</script>"
    )
    with session_factory() as session:
        job = session.get(JobPosting, job_id)
        assert job is not None
        job.snapshots.append(
            JobSnapshot(
                content_hash="raw-html",
                raw_html_compressed=gzip.compress(
                    json.dumps(
                        {
                            "description": raw_description,
                            "description_format": "html",
                            "metadata": {},
                            "raw": {},
                        }
                    ).encode()
                ),
                cleaned_text=job.description_text,
            )
        )
        session.commit()

    class FakeSummaryProvider:
        def __init__(self) -> None:
            self.calls: list[tuple[object, object, dict[str, str]]] = []

        async def summarize_job(self, job, profile, **options):
            self.calls.append((job, profile, options))
            return CodexJobSummary(
                overview="Direkte Unterstützung der Geschäftsführung in Stuttgart.",
                key_points=(
                    "Vollzeit und unbefristet",
                    "Präsenz in Stuttgart",
                    "Termin- und Reiseplanung",
                ),
                missing_information=("Vergütung fehlt",),
            )

    provider = FakeSummaryProvider()
    client.app.state.codex_provider = provider

    pending = client.get(f"/jobs/{job_id}")
    assert pending.status_code == 200
    assert 'data-summary-state="pending"' in pending.text
    assert "<summary><span>Originaltext anzeigen</span>" in pending.text
    assert "<h3>Deine Aufgaben</h3>" in pending.text
    assert "<li>Termine koordinieren</li><li>Reisen planen</li>" in pending.text
    assert "alert('nicht ausführen')" not in pending.text

    generated = client.post(f"/jobs/{job_id}/summary")
    assert generated.status_code == 200
    assert generated.json()["overview"].startswith("Direkte Unterstützung")
    assert generated.json()["key_points"][0] == "Vollzeit und unbefristet"
    assert provider.calls[0][2] == {
        "locale": "de",
        "model": "gpt-5.6-luna",
        "reasoning_effort": "medium",
    }
    assert "Anna Beispiel" not in repr(provider.calls[0][1])

    cached = client.post(f"/jobs/{job_id}/summary")
    assert cached.status_code == 200
    assert len(provider.calls) == 1
    ready = client.get(f"/jobs/{job_id}")
    assert 'data-summary-state="ready"' in ready.text
    assert "Direkte Unterstützung der Geschäftsführung in Stuttgart." in ready.text

    with session_factory() as session:
        preference = session.get(PreferenceProfile, 1)
        assert preference is not None
        preference.version += 1
        session.commit()
    stale = client.get(f"/jobs/{job_id}")
    assert 'data-summary-state="pending"' in stale.text
    refreshed = client.post(f"/jobs/{job_id}/summary")
    assert refreshed.status_code == 200
    assert len(provider.calls) == 2


def test_profile_sources_settings_and_pipeline_render(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _seed_jobs(session_factory)

    profile = client.get("/profile")
    assert profile.status_code == 200
    assert "Erfahrung und Fähigkeiten" in profile.text
    assert "Präsenz bevorzugt" in profile.text
    assert 'name="preference_prompt"' in profile.text
    assert "Arbeitszeit: Vollzeit" in profile.text
    assert "Vertragsarten: Unbefristet" in profile.text
    assert "Bevorzugte Regionen: Baden-Württemberg" in profile.text
    assert "Deutschlandweite Remote-Stellen einbeziehen: Ja" in profile.text
    assert "Sprachen: Deutsch, Englisch" in profile.text
    assert 'name="role_taxonomy"' in profile.text
    assert "Hauptbegriffe:" in profile.text
    assert "Assistenz der Geschäftsführung" in profile.text
    assert "Ähnliche Begriffe mit KI generieren" in profile.text
    assert "Zusätzliche Rollenbegriffe:" not in profile.text
    warning_position = profile.text.index("Änderungen wirken sich auf deine Ergebnisse aus")
    editor_position = profile.text.index('id="role_taxonomy"')
    assert warning_position < editor_position
    assert 'data-preference-undo' in profile.text
    assert 'data-role-taxonomy-reset' in profile.text
    assert 'id="approved_preference_prompt"' in profile.text
    assert 'id="approved_role_taxonomy"' in profile.text
    assert "Weiter: prüfen und speichern" in profile.text
    assert "Letzte Eingabe rückgängig machen" in profile.text
    assert 'name="workplace_preference"' not in profile.text
    assert 'name="commute_soft"' not in profile.text
    assert 'name="commute_max"' not in profile.text

    sources = client.get("/sources")
    assert sources.status_code == 200
    assert "Bundesagentur für Arbeit" in sources.text
    assert "Kostenlose API" in sources.text
    assert "Gespeicherte Stellen" in sources.text
    assert "<strong>3</strong> Stellen" in sources.text
    assert "API-Abrufe sind deaktiviert" in sources.text

    settings = client.get("/settings")
    assert settings.status_code == 200
    assert "Noch nicht verbunden" in settings.text
    assert "Automatische API-Abrufe" in settings.text
    assert "Firecrawl-Crawling" in settings.text
    assert "Standardansicht" not in settings.text
    assert 'name="fulltime_default"' not in settings.text
    assert 'name="permanent_default"' not in settings.text
    assert "Stellenfilter stehen ausschließlich in der Stellensuche" in settings.text

    applications = client.get("/applications")
    assert applications.status_code == 200
    assert "Merkliste" in applications.text
    assert "Es wird nichts versendet" in applications.text



def test_planned_status_requires_confirmation_when_availability_is_unverified(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    job_id = _seed_jobs(session_factory)[0]

    blocked = client.post(
        f"/jobs/{job_id}/status",
        data={"status": "Bewerbung geplant", "notes": "Heute vorbereiten"},
    )

    assert blocked.status_code == 409
    assert "Prüfung nicht möglich" in blocked.text
    with session_factory() as session:
        application = session.query(ApplicationRecord).one()
        assert application.status == "Merkliste"

    confirmed = client.post(
        f"/jobs/{job_id}/status",
        data={
            "status": "Bewerbung geplant",
            "notes": "Heute vorbereiten",
            "confirm_unverified": "true",
        },
        follow_redirects=False,
    )

    assert confirmed.status_code == 303
    with session_factory() as session:
        application = session.query(ApplicationRecord).one()
        assert application.status == "Bewerbung geplant"
        assert application.notes == "Heute vorbereiten"


def test_application_drafts_require_confirmation_when_availability_check_is_blocked(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    job_ids = _seed_jobs(session_factory)
    with session_factory() as session:
        profile = session.query(CandidateProfile).one()
        profile.is_confirmed = True
        profile.structured_data = {**profile.structured_data, "skills": ["Kalendersteuerung"]}
        session.commit()
    client.app.state.codex_provider = WebDraftProvider()

    response = client.post(
        "/application-drafts/generate",
        data={"mode": "selected", "job_ids": str(job_ids[0])},
    )

    assert response.status_code == 409
    assert "Prüfung nicht möglich" in response.text
    assert "Trotz fehlender Bestätigung fortfahren" in response.text
    with session_factory() as session:
        assert session.query(ApplicationDraft).count() == 0


def test_application_drafts_generate_render_and_create_edit_revision(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    job_ids = _seed_jobs(session_factory)
    with session_factory() as session:
        profile = session.query(CandidateProfile).one()
        profile.is_confirmed = True
        profile.structured_data = {
            **profile.structured_data,
            "skills": ["Kalendersteuerung"],
        }
        session.commit()
    client.app.state.codex_provider = WebDraftProvider()

    generated = client.post(
        "/application-drafts/generate",
        data={
            "mode": "selected",
            "job_ids": str(job_ids[0]),
            "confirm_unverified": "true",
        },
        follow_redirects=False,
    )

    assert generated.status_code == 303
    assert generated.headers["location"].endswith("/application-drafts/1")
    detail = client.get(generated.headers["location"])
    assert detail.status_code == 200
    assert "Lebenslauf mit Kalendersteuerung" in detail.text
    assert "Freigegeben." in detail.text
    assert "Kein Versand" in detail.text

    edited = client.post(
        generated.headers["location"],
        data={
            "cv_draft": "Manuell überarbeiteter Lebenslauf",
            "cover_letter": "Manuell überarbeitetes Anschreiben",
        },
        follow_redirects=False,
    )
    assert edited.status_code == 303
    assert edited.headers["location"].endswith("/application-drafts/2")
    edited_detail = client.get(edited.headers["location"])
    assert "Manuell überarbeiteter Lebenslauf" in edited_detail.text
    assert "Erneute HR-Prüfung erforderlich" in edited_detail.text

    with session_factory() as session:
        drafts = session.query(ApplicationDraft).order_by(ApplicationDraft.revision).all()
        assert [draft.revision for draft in drafts] == [1, 2]
        assert drafts[0].cv_draft == "Anna Beispiel\n\nLebenslauf mit Kalendersteuerung"


def test_settings_hides_raw_codex_status_errors_and_offers_reconnect(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FailingCodexProvider:
        async def status(self) -> dict[str, object]:
            raise RuntimeError("401 Unauthorized: token_invalidated")

    monkeypatch.setattr(
        "jobradar.web.get_settings",
        lambda: Settings(codex_enabled=True),
    )
    client.app.state.codex_provider = FailingCodexProvider()

    response = client.get("/language/en", params={"next": "/settings"})

    assert response.status_code == 200
    assert "The saved ChatGPT sign-in has expired or is currently unavailable." in response.text
    assert "Reconnect ChatGPT" in response.text
    assert "401 Unauthorized" not in response.text
    assert "token_invalidated" not in response.text


def test_settings_hides_device_code_until_disclosure(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class DeviceLoginProvider:
        async def status(self) -> dict[str, object]:
            return {"enabled": True, "connected": False}

        async def begin_login(self) -> dict[str, str]:
            return {
                "verificationUrl": "https://example.invalid/device",
                "userCode": "TEST-CODE",
            }

    monkeypatch.setattr(
        "jobradar.web.get_settings",
        lambda: Settings(codex_enabled=True),
    )
    client.app.state.codex_provider = DeviceLoginProvider()
    client.get("/language/en", params={"next": "/settings"})

    response = client.post("/settings/codex-login")

    assert response.status_code == 200
    marker = '<details class="codex-code-disclosure">'
    assert marker in response.text
    disclosure = response.text.split(marker, maxsplit=1)[1].split("</details>", maxsplit=1)[0]
    assert "Show connection code" in disclosure
    assert "Hide connection code" in disclosure
    assert "TEST-CODE" in disclosure


def test_sources_distinguish_enabled_but_unchecked_crawler(
    client: TestClient,
    session_factory: sessionmaker[Session],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime = Settings(
        crawling_enabled=True,
        ba_jobs_enabled=True,
        firecrawl_enabled=True,
    )
    monkeypatch.setattr("jobradar.web.get_settings", lambda: runtime)
    monkeypatch.setattr(setting_service, "get_runtime_settings", lambda: runtime)
    with session_factory() as session:
        session.add_all(
            [
                AppSetting(
                    key=setting_service.CRAWLING_SETTING_KEY,
                    value_json=True,
                ),
                AppSetting(
                    key=setting_service.FIRECRAWL_SETTING_KEY,
                    value_json=True,
                ),
            ]
        )
        session.add_all(
            [
                Source(
                    name="Bundesagentur für Arbeit – Jobsuche",
                    slug="ba_jobsuche",
                    enabled=True,
                    status="healthy",
                ),
                Source(
                    name="Firecrawl (selbst gehostet)",
                    slug="firecrawl_self_hosted",
                    kind="crawler",
                    enabled=True,
                    status="prepared",
                ),
            ]
        )
        session.commit()

    response = client.get("/sources")

    assert response.status_code == 200
    assert response.text.count(">Aktiv</span>") == 1
    assert response.text.count(">Noch nicht geprüft</span>") == 2
    assert ">Pausiert</span>" not in response.text
    assert "noch nicht erfolgreich ausgeführt" in response.text
    assert "Automatische API-Abrufe aktiv" in response.text

def test_sources_show_failed_crawler_as_error(
    client: TestClient,
    session_factory: sessionmaker[Session],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime = Settings(crawling_enabled=True, firecrawl_enabled=True)
    monkeypatch.setattr("jobradar.web.get_settings", lambda: runtime)
    monkeypatch.setattr(setting_service, "get_runtime_settings", lambda: runtime)
    with session_factory() as session:
        session.add_all(
            [
                AppSetting(
                    key=setting_service.CRAWLING_SETTING_KEY,
                    value_json=True,
                ),
                AppSetting(
                    key=setting_service.FIRECRAWL_SETTING_KEY,
                    value_json=True,
                ),
            ]
        )
        source = Source(
            name="Firecrawl (selbst gehostet)",
            slug="firecrawl_self_hosted",
            kind="crawler",
            enabled=True,
            status="error",
            last_error="ConnectError: connection refused",
            metadata_json={
                "firecrawl": {
                    "targets": [
                        {
                            "url": "https://portal.example/jobs",
                            "label": "Portal",
                            "target_kind": "job_portal",
                        }
                    ]
                }
            },
        )
        session.add(source)
        session.flush()
        failed_at = datetime.now(UTC)
        session.add(
            CrawlRun(
                source_id=source.id,
                status="failed",
                started_at=failed_at,
                finished_at=failed_at,
                error_count=1,
                error_message="ConnectError: connection refused",
                details={
                    "connector": {
                        "target_results": [
                            {
                                "target": "https://portal.example/jobs",
                                "status": "failed",
                            }
                        ]
                    }
                },
            )
        )
        session.commit()

    response = client.get("/sources?crawl_tab=true")

    assert response.status_code == 200
    assert response.text.count(">Fehler</span>") == 2
    assert ">Aktiv</span>" not in response.text
    assert "Letzter Crawlerlauf fehlgeschlagen." in response.text
    assert "ConnectError: connection refused" in response.text

def test_sources_replace_empty_target_status_with_exact_page_error(
    client: TestClient,
    session_factory: sessionmaker[Session],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime = Settings(crawling_enabled=True, firecrawl_enabled=True)
    monkeypatch.setattr("jobradar.web.get_settings", lambda: runtime)
    monkeypatch.setattr(setting_service, "get_runtime_settings", lambda: runtime)
    target_url = "https://portal.example/jobs"
    first_document_url = "https://portal.example/jobs/executive-assistant"
    second_document_url = "https://portal.example/jobs/office-manager"
    hidden_document_url = "https://portal.example/jobs/chief-of-staff"
    exact_error = "HTTPStatusError: Firecrawl returned 500 Internal Server Error"
    with session_factory() as session:
        session.add_all(
            [
                AppSetting(key=setting_service.CRAWLING_SETTING_KEY, value_json=True),
                AppSetting(key=setting_service.FIRECRAWL_SETTING_KEY, value_json=True),
            ]
        )
        source = Source(
            name="Firecrawl (selbst gehostet)",
            slug="firecrawl_self_hosted",
            kind="crawler",
            enabled=True,
            status="error",
            metadata_json={
                "firecrawl": {
                    "targets": [
                        {
                            "url": target_url,
                            "label": "Portal",
                            "target_kind": "job_portal",
                        }
                    ]
                }
            },
        )
        session.add(source)
        session.flush()
        finished_at = datetime.now(UTC)
        session.add(
            CrawlRun(
                source_id=source.id,
                status="partial",
                started_at=finished_at,
                finished_at=finished_at,
                error_count=3,
                error_message=exact_error,
                details={
                    "connector": {
                        "target_results": [
                            {
                                "target": target_url,
                                "status": "empty",
                                "documents_seen": 0,
                                "records_emitted": 0,
                                "pages_rejected": 2,
                                "duplicates_skipped": 1,
                                "query_filtered": 3,
                                "page_error_count": 3,
                            }
                        ],
                        "page_errors": [
                            {
                                "target": target_url,
                                "document_url": first_document_url,
                                "error": exact_error,
                            },
                            {
                                "target": target_url,
                                "document_url": second_document_url,
                                "error": exact_error,
                            },
                            {
                                "target": target_url,
                                "document_url": hidden_document_url,
                                "error": exact_error,
                            },
                        ],
                    }
                },
            )
        )
        session.commit()

    response = client.get("/sources?crawl_tab=true")

    assert response.status_code == 200
    assert ">Technischer Fehler</span>" in response.text
    assert ">Ohne Treffer</span>" not in response.text
    assert "Letzter Crawlerlauf teilweise fehlgeschlagen." in response.text
    assert "Dokumente geprüft: 0" in response.text
    assert "Seiten verworfen: 2" in response.text
    assert "Technische Fehler: 3" in response.text
    assert '<details class="crawl-target-errors">' in response.text
    assert "Exakte technische Fehler <span>3</span>" in response.text
    assert first_document_url in response.text
    assert second_document_url in response.text
    assert hidden_document_url not in response.text
    assert exact_error in response.text



def test_profile_adds_and_persists_multiple_start_addresses(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    profile_page = client.get("/profile")
    assert profile_page.status_code == 200
    assert 'name="addresses"' in profile_page.text
    assert "data-address-add" in profile_page.text
    assert 'name="cv"' in profile_page.text
    assert 'data-cv-state="empty"' in profile_page.text
    assert "Datei auswählen" in profile_page.text
    assert "Erneut hochladen" not in profile_page.text

    saved = client.post(
        "/profile",
        data={
            "full_name": "Anna Beispiel",
            "addresses": [
                "  Musterstraße 1  ",
                "Zweitweg 2",
                "MUSTERSTRASSE 1",
            ],
            "languages": "Deutsch, Englisch",
        },
        files={
            "cv": (
                "Lebenslauf.docx",
                _cv_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert saved.status_code == 303

    with session_factory() as session:
        candidate = session.query(CandidateProfile).one()
        addresses = session.query(CandidateAddress).order_by(CandidateAddress.position).all()
        assert candidate.home_location == "Musterstraße 1"
        assert candidate.cv_filename == "Lebenslauf.docx"
        assert candidate.structured_data["languages"] == ["Deutsch", "Englisch"]
        assert [address.address_text for address in addresses] == [
            "Musterstraße 1",
            "Zweitweg 2",
        ]

    rendered = client.get("/profile")
    assert 'value="Musterstraße 1"' in rendered.text
    assert 'value="Zweitweg 2"' in rendered.text
    assert 'value="None"' not in rendered.text
    assert 'data-cv-state="accepted"' in rendered.text
    assert "Akzeptiert" in rendered.text
    assert "Der Lebenslauf wurde erfolgreich eingelesen." in rendered.text
    assert "Erneut hochladen" in rendered.text

    removed = client.post(
        "/profile",
        data={
            "addresses": ["Musterstraße 1"],
        },
        follow_redirects=False,
    )
    assert removed.status_code == 303
    with session_factory() as session:
        addresses = session.query(CandidateAddress).order_by(CandidateAddress.position).all()
        assert [address.address_text for address in addresses] == ["Musterstraße 1"]


def test_profile_save_does_not_override_textbox_preferences(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _seed_jobs(session_factory)

    saved = client.post(
        "/profile",
        data={
            "full_name": "Anna Beispiel",
            "languages": "Französisch, Englisch",
            "preference_prompt": "Bevorzugte Regionen: Bayern",
            "workplace_preference": "remote",
            "commute_soft": "10",
            "commute_max": "20",
        },
        follow_redirects=False,
    )

    assert saved.status_code == 303
    with session_factory() as session:
        candidate = session.query(CandidateProfile).one()
        preferences = session.query(PreferenceProfile).one()
        assert candidate.structured_data["languages"] == ["Französisch", "Englisch"]
        assert preferences.languages == ["Deutsch", "Englisch"]
        assert preferences.preferred_states == ["Baden-Württemberg"]
        assert preferences.onsite_weight == 1.0
        assert preferences.commute_penalty_minutes == 60


def test_profile_groups_duplicate_cv_evidence(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    with session_factory() as session:
        session.add(
            CandidateProfile(
                structured_data={
                    "cv_import": {
                        "evidence": [
                            {
                                "field": "skills",
                                "value": "Excel",
                                "source_text": "Skills: Excel, PowerPoint",
                            },
                            {
                                "field": "skills",
                                "value": "PowerPoint",
                                "source_text": "Skills: Excel, PowerPoint",
                            },
                            {
                                "field": "skills",
                                "value": "Excel",
                                "source_text": "Skills: Excel, PowerPoint",
                            },
                            {
                                "field": "summary",
                                "value": "Experienced executive support professional.",
                                "source_text": "Experienced executive support professional.",
                            },
                        ]
                    }
                }
            )
        )
        session.commit()

    rendered = client.get("/profile")

    assert rendered.status_code == 200
    assert "Wichtige Fähigkeiten: Excel, PowerPoint" in rendered.text
    assert rendered.text.count("Skills: Excel, PowerPoint") == 1
    assert rendered.text.count("Experienced executive support professional.") == 1


def test_profile_reviews_prompt_without_applying_settings(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _seed_jobs(session_factory)

    reviewed = client.post(
        "/profile/preferences/review",
        data={
            "preference_prompt": "Bevorzugte Regionen: Bayern",
            "preference_editor": "preference_prompt",
        },
    )

    assert reviewed.status_code == 200
    assert "Prüfen und speichern" in reviewed.text
    assert "Geprüfte Änderungen speichern" in reviewed.text
    assert 'data-preference-review tabindex="-1"' in reviewed.text
    assert "Noch nicht gespeichert" in reviewed.text
    assert "data-preference-reset" in reviewed.text
    assert "Auf gespeicherten Stand zurücksetzen" in reviewed.text
    assert "Baden-Württemberg" in reviewed.text
    assert "Bayern" in reviewed.text
    assert "Neue Suche erforderlich" in reviewed.text
    assert "Crawling ist derzeit gesperrt" in reviewed.text
    assert reviewed.text.index('id="preference_prompt"') < reviewed.text.index(
        'id="preference-review-title"'
    )
    approved_start = reviewed.text.index('id="approved_preference_prompt"')
    approved_end = reviewed.text.index("</textarea>", approved_start)
    approved_prompt = reviewed.text[approved_start:approved_end]
    assert "Baden-Württemberg" in approved_prompt
    assert "Bayern" not in approved_prompt
    with session_factory() as session:
        preferences = session.query(PreferenceProfile).one()
        assert preferences.preferred_states == ["Baden-Württemberg"]



def test_profile_shows_prompt_errors_beside_the_edited_textarea(client: TestClient) -> None:
    reviewed = client.post(
        "/profile/preferences/review",
        data={
            "preference_prompt": "Diese Anweisung gibt es nicht",
            "preference_editor": "preference_prompt",
        },
    )

    assert reviewed.status_code == 200
    textarea_position = reviewed.text.index('id="preference_prompt"')
    error_position = reviewed.text.index('id="preference-prompt-errors"')
    assert textarea_position < error_position
    assert 'aria-describedby="preference-prompt-help preference-prompt-errors"' in reviewed.text
    assert 'aria-invalid="true"' in reviewed.text
    assert "Diese Zeile wurde nicht verstanden" in reviewed.text
    assert 'id="preference-review-title"' not in reviewed.text

def test_profile_generates_role_titles_without_applying_them(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    _seed_jobs(session_factory)

    class RoleProvider:
        calls: list[tuple[list[str], list[str], str]] = []

        async def generate_similar_role_titles(
            self,
            primary_terms: list[str],
            existing_titles: list[str],
            *,
            locale: str,
        ) -> tuple[GeneratedRoleTermGroup, ...]:
            self.calls.append((primary_terms, existing_titles, locale))
            return (
                GeneratedRoleTermGroup(
                    primary_term="Assistenz der Geschäftsführung",
                    titles=(
                        "Geschäftsführungsassistenz",
                        "Assistenz der Geschäftsführung",
                    ),
                ),
            )

    provider = RoleProvider()
    client.app.state.codex_provider = provider
    taxonomy = """Hauptbegriffe:
- Assistenz der Geschäftsführung

Zusätzliche Titel:"""

    generated = client.post(
        "/profile/preferences/roles/generate",
        data={"preference_prompt": "", "role_taxonomy": taxonomy},
    )

    assert generated.status_code == 200
    assert "Geschäftsführungsassistenz" in generated.text
    assert "1 ähnliche Rollentitel wurden ergänzt" in generated.text
    assert provider.calls == [
        (["Assistenz der Geschäftsführung"], [], "de")
    ]
    with session_factory() as session:
        preferences = session.query(PreferenceProfile).one()
        assert preferences.additional_role_terms == []


def test_profile_applies_editable_role_taxonomy(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    _seed_jobs(session_factory)
    taxonomy = """Hauptbegriffe:
- Executive Assistant
- Chief of Staff

Zusätzliche Titel:
- Executive Office Coordinator"""

    applied = client.post(
        "/profile/preferences/apply",
        data={
            "preference_prompt": "",
            "role_taxonomy": taxonomy,
            "preference_action": "apply",
        },
        follow_redirects=False,
    )

    assert applied.status_code == 303
    assert "coverage_pending=true" in applied.headers["location"]
    with session_factory() as session:
        preferences = session.query(PreferenceProfile).one()
        assert preferences.primary_role_terms == ["Executive Assistant", "Chief of Staff"]
        assert preferences.additional_role_terms == ["Executive Office Coordinator"]


def test_profile_applies_prompt_updates_results_and_can_undo(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _seed_jobs(session_factory)

    applied = client.post(
        "/profile/preferences/apply",
        data={
            "preference_prompt": "Bevorzugte Regionen: Bayern",
            "preference_action": "apply",
        },
        follow_redirects=False,
    )

    assert applied.status_code == 303
    assert "preferences_updated=true" in applied.headers["location"]
    assert "coverage_pending=true" in applied.headers["location"]
    rendered = client.get("/profile")
    prompt_start = rendered.text.index(">", rendered.text.index('id="preference_prompt"')) + 1
    prompt_end = rendered.text.index("</textarea>", prompt_start)
    assert rendered.text[prompt_start:prompt_end] == "Bevorzugte Regionen: Bayern"
    with session_factory() as session:
        preferences = session.query(PreferenceProfile).one()
        assert preferences.preferred_states == ["Bayern"]
        history = session.get(AppSetting, "preferences.last_applied_snapshot")
        assert history is not None
        assert history.value_json["preferred_states"] == ["Baden-Württemberg"]

    matches = client.get("/jobs", params={"view": "matches"})
    assert "Muster GmbH" not in matches.text
    assert "Remote SE" in matches.text

    undone = client.post("/profile/preferences/undo", follow_redirects=False)

    assert undone.status_code == 303
    with session_factory() as session:
        preferences = session.query(PreferenceProfile).one()
        assert preferences.preferred_states == ["Baden-Württemberg"]


def test_profile_prompt_refresh_obeys_crawling_lock(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _seed_jobs(session_factory)

    applied = client.post(
        "/profile/preferences/apply",
        data={
            "preference_prompt": "Bevorzugte Regionen: Bayern",
            "preference_action": "apply_refresh",
        },
        follow_redirects=False,
    )

    assert applied.status_code == 303
    assert "crawl_blocked=true" in applied.headers["location"]
    with session_factory() as session:
        crawl = session.query(CrawlRun).one()
        assert crawl.status == "blocked"
        assert crawl.details["network_called"] is False


def test_missing_job_has_helpful_404(client: TestClient) -> None:
    response = client.get("/jobs/999999")
    assert response.status_code == 404
    assert "Diese Stelle wurde nicht gefunden" in response.text


def test_job_detail_hides_unsafe_provider_link(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    job_id = _seed_jobs(session_factory)[0]
    with session_factory() as session:
        job = session.get(JobPosting, job_id)
        assert job is not None
        job.canonical_url = "javascript:alert(document.cookie)"
        session.commit()

    response = client.get(f"/jobs/{job_id}")

    assert response.status_code == 200
    assert "javascript:" not in response.text


def test_language_switch_translates_the_complete_interface(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    job_id = _seed_jobs(session_factory)[0]

    selected = client.get(
        "/language/en", params={"next": f"/jobs/{job_id}"}, follow_redirects=False
    )
    assert selected.status_code == 303
    assert selected.headers["location"].endswith(f"/jobs/{job_id}")


    jobs = client.get("/jobs")
    assert "Best matches" in jobs.text
    assert "Role relevance" in jobs.text
    assert "Employer" in jobs.text
    assert "On-site in Baden-Württemberg" in jobs.text
    assert "Remote throughout Germany" in jobs.text

    detail = client.get(f"/jobs/{job_id}")
    assert "Role relevance" in detail.text
    assert "Candidate fit" in detail.text
    assert "Direct support for executive management" in detail.text
    assert "Application status" in detail.text
    assert "Saved" in detail.text

    english_profile = client.get("/profile")
    assert "My profile" in english_profile.text
    assert "Starting addresses for commute times" in english_profile.text
    assert "Add address" in english_profile.text
    assert "Source status" in client.get("/sources").text
    settings = client.get("/settings")
    assert "Automatic API retrieval" in settings.text
    assert "Firecrawl crawling" in settings.text
    assert "Nothing is submitted." in client.get("/applications").text

    restored = client.get("/language/de", params={"next": "/"})
    assert restored.status_code == 200
    assert '<html lang="de">' in restored.text
    assert "Jobportale: nicht geprüft" in restored.text
    assert "Firmenportale: nicht geprüft" in restored.text
